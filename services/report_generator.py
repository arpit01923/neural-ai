from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def build_markdown_report(report: dict) -> str:
    title = report.get("title", "Research Report")
    summary = report.get("summary", "")
    sections = report.get("sections", [])
    sources = report.get("sources", [])

    lines = [f"# {title}", "", "## Overview", "", summary, ""]

    for section in sections:
        heading = section.get("heading", "Section")
        content = section.get("content", "")
        lines.extend([f"## {heading}", "", content, ""])

    lines.append("## References")
    lines.append("")
    for index, source in enumerate(sources, start=1):
        lines.append(f"{index}. [{source.get('title', 'Untitled')}]({source.get('url', '')})")

    return "\n".join(lines).strip() + "\n"


def build_pdf_bytes(report: dict) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    flow = [Paragraph(report.get("title", "Research Report"), styles["Title"]), Spacer(1, 12)]

    flow.append(Paragraph(report.get("summary", ""), styles["BodyText"]))
    flow.append(Spacer(1, 12))

    for section in report.get("sections", []):
        flow.append(Paragraph(section.get("heading", "Section"), styles["Heading2"]))
        flow.append(Paragraph(section.get("content", ""), styles["BodyText"]))
        flow.append(Spacer(1, 8))

    doc.build(flow)
    return buffer.getvalue()


def build_docx_bytes(report: dict) -> bytes:
    document = Document()
    document.add_heading(report.get("title", "Research Report"), level=1)
    document.add_paragraph(report.get("summary", ""))

    for section in report.get("sections", []):
        document.add_heading(section.get("heading", "Section"), level=2)
        document.add_paragraph(section.get("content", ""))

    document.add_heading("References", level=2)
    for source in report.get("sources", []):
        document.add_paragraph(f"- {source.get('title', 'Untitled')}: {source.get('url', '')}")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
